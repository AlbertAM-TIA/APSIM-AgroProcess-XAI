"""Build harmonized annual APSIM records for ANOVA and SHAP.

The script reads:
  * two individual-practice DailyReport CSV files; and
  * six combined-system DailyReport worksheets.

It preserves the raw files and writes annual, analysis-ready CSV outputs.
Crop yields are annual maximum harvest yields. Crop-specific stress and
resource variables are summarized only across days with positive crop biomass.
Annual SOC change is the within-calendar-year difference between the last and
first daily 0-30 cm SOC stock values.
"""

from __future__ import annotations

import argparse
import gc
from pathlib import Path
import os
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(os.environ.get("RA_PROJECT_ROOT", Path(__file__).resolve().parents[2]))
INDIVIDUAL_DIR = ROOT / "data/raw/individual"
COMBINED_DIR = ROOT / "data/raw/combined"
OUTPUT_DIR = ROOT / "outputs/anova_shap_raw"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
INTERMEDIATE_DIR = OUTPUT_DIR / "intermediate"
INTERMEDIATE_DIR.mkdir(parents=True, exist_ok=True)


BASE_COLUMNS = [
    "SimulationName",
    "Zone",
    "Date",
    "RotationManager.CurrentState",
    "SOC_Accrual_30cm_tonnes_C_ha",
    "SoilEvaporation",
    "TotalSoilN",
    "TotalSoilWaterasTotalSoilWater",
    "NitrogenApplied",
    "OM_residue",
    "PAWC",
]

CROP_SPECS = {
    "Wheat": {
        "yield": "WheatYieldkgha",
        "biomass": "WheatBiomass",
        "lai": "Wheat_LAIm2m",
        "co2": "Wheat_CO2_stress",
        "nitrogen": "Wheat_nitrogen_stress",
        "temperature": "Wheat_temp_stress",
        "vpd": "Wheat_VPD_stress",
        "water": "Wheat_water_stress",
    },
    "Barley": {
        "yield": "BarleyYieldkgha",
        "biomass": "BarleyBiomass",
        "lai": "baley_LAIm2m",
        "co2": "barley_CO2_stress",
        "nitrogen": "barley_nitrogen_stress",
        "temperature": "barley_temp_stress",
        "vpd": "barley_VPD_stress",
        "water": "barley_water_stress",
    },
    "Rapeseed oil": {
        "yield": "CanolaYieldkgha",
        "biomass": "CanolaBiomass",
        "lai": "canola_LAIm2m",
        "co2": "canola_CO2_stress",
        "nitrogen": "canola_nitrogen_stress",
        "temperature": "canola_temp_stress",
        "vpd": "canola_VPD_stress",
        "water": "canola_water_stress",
    },
}

SELECTED_COLUMNS = list(
    dict.fromkeys(
        BASE_COLUMNS
        + [
            col
            for spec in CROP_SPECS.values()
            for col in spec.values()
        ]
    )
)


def co2_metadata(simulation_name: str, period: str) -> tuple[str, float]:
    text = str(simulation_name)
    if "FromFile" in text:
        return "Trend", np.nan

    match = re.search(r"Value(\d+(?:\.\d+)?)", text)
    if not match:
        return "Unknown", np.nan

    ppm = float(match.group(1))
    lower = 390.0 if period == "Historical" else 510.0
    higher = 420.0 if period == "Historical" else 570.0

    if np.isclose(ppm, lower):
        return "Lower fixed", ppm
    if np.isclose(ppm, higher):
        return "Higher fixed", ppm
    return "Other fixed", ppm


def rotation_from_zone(zone: str) -> str:
    return "Spring" if "Winter_and_Spring" in str(zone) else "Winter"


def individual_system(simulation_name: str, zone: str) -> str:
    text = f"{simulation_name} {zone}".lower()
    if "baseline" in text:
        return "Baseline"
    if "cover crop" in text:
        return "Simple cover crop"
    if "crop diversity" in text or "diversity" in text:
        return "Crop diversity"
    if "herbal" in text:
        return "Herbal leys"
    if "livestock" in text:
        return "Integrate livestock"
    if "lown_10" in text:
        return "Fertiliser reduction (10%)"
    if "mediumn_30" in text:
        return "Fertiliser reduction (30%)"
    if "highn_50" in text:
        return "Fertiliser reduction (50%)"
    return "Unclassified"


def to_numeric_frame(data: pd.DataFrame) -> pd.DataFrame:
    out = data.copy()
    out["Date"] = pd.to_datetime(out["Date"], errors="coerce")
    numeric_columns = [
        col
        for col in SELECTED_COLUMNS
        if col not in {
            "SimulationName",
            "Zone",
            "Date",
            "RotationManager.CurrentState",
        }
    ]
    for col in numeric_columns:
        if col not in out:
            out[col] = np.nan
        out[col] = pd.to_numeric(out[col], errors="coerce")
    return out


def summarize_one_source(
    data: pd.DataFrame,
    *,
    period: str,
    scope: str,
    system_override: str | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    data = to_numeric_frame(data)
    data = data.loc[data["Date"].notna()].copy()
    data["Year"] = data["Date"].dt.year

    annual_soc_rows: list[dict] = []
    crop_rows: list[dict] = []

    for (simulation_name, zone, year), group in data.groupby(
        ["SimulationName", "Zone", "Year"],
        observed=True,
        sort=False,
    ):
        group = group.sort_values("Date")
        co2_scenario, co2_ppm = co2_metadata(simulation_name, period)
        rotation = rotation_from_zone(zone)
        system = (
            system_override
            if system_override is not None
            else individual_system(simulation_name, zone)
        )

        soc_values = group["SOC_Accrual_30cm_tonnes_C_ha"].dropna()
        soc_start = soc_values.iloc[0] if len(soc_values) else np.nan
        soc_end = soc_values.iloc[-1] if len(soc_values) else np.nan

        common = {
            "Scope": scope,
            "System": system,
            "Period": period,
            "Rotation": rotation,
            "CO2_scenario": co2_scenario,
            "CO2_ppm": co2_ppm,
            "Year": int(year),
            "Year_index": int(year - group["Year"].min() + 1),
            "SimulationName": simulation_name,
            "Zone": zone,
            "n_daily_rows": len(group),
        }

        annual_soc_rows.append(
            {
                **common,
                "SOC_start_t_C_ha": soc_start,
                "SOC_end_t_C_ha": soc_end,
                "SOC_change_t_C_ha_yr": soc_end - soc_start,
                "Soil_water_mean": group[
                    "TotalSoilWaterasTotalSoilWater"
                ].mean(),
                "Soil_water_min": group[
                    "TotalSoilWaterasTotalSoilWater"
                ].min(),
                "Soil_N_mean": group["TotalSoilN"].mean(),
                "Soil_N_min": group["TotalSoilN"].min(),
                "Nitrogen_applied_kg_ha": group["NitrogenApplied"].sum(
                    min_count=1
                ),
                "Soil_evaporation_sum": group["SoilEvaporation"].sum(
                    min_count=1
                ),
                "PAWC_mean": group["PAWC"].mean(),
                "OM_residue_mean": group["OM_residue"].mean(),
                "OM_residue_start": group["OM_residue"].dropna().iloc[0]
                if group["OM_residue"].notna().any()
                else np.nan,
                "OM_residue_end": group["OM_residue"].dropna().iloc[-1]
                if group["OM_residue"].notna().any()
                else np.nan,
                "Total_crop_biomass_max": sum(
                    group[spec["biomass"]].max(skipna=True)
                    for spec in CROP_SPECS.values()
                    if spec["biomass"] in group
                ),
                "Total_crop_yield_kg_ha": sum(
                    max(group[spec["yield"]].max(skipna=True), 0)
                    for spec in CROP_SPECS.values()
                    if spec["yield"] in group
                ),
            }
        )

        for crop, spec in CROP_SPECS.items():
            yield_values = group[spec["yield"]]
            max_yield = yield_values.max(skipna=True)
            if not np.isfinite(max_yield) or max_yield <= 0:
                continue

            active = group[spec["biomass"]].fillna(0) > 0
            active_group = group.loc[active]
            if active_group.empty:
                active_group = group.loc[yield_values.fillna(0) > 0]
            if active_group.empty:
                continue

            max_yield_index = yield_values.idxmax()
            harvest_date = group.loc[max_yield_index, "Date"]

            crop_rows.append(
                {
                    **common,
                    "Crop": crop,
                    "Yield_t_ha": max_yield / 1000.0,
                    "Biomass_max_kg_ha": active_group[
                        spec["biomass"]
                    ].max(skipna=True),
                    "LAI_max": active_group[spec["lai"]].max(skipna=True),
                    "CO2_factor_mean": active_group[spec["co2"]].mean(),
                    "N_stress_factor_mean": active_group[
                        spec["nitrogen"]
                    ].mean(),
                    "N_stress_factor_min": active_group[
                        spec["nitrogen"]
                    ].min(),
                    "Temperature_factor_mean": active_group[
                        spec["temperature"]
                    ].mean(),
                    "Temperature_factor_min": active_group[
                        spec["temperature"]
                    ].min(),
                    "VPD_factor_mean": active_group[spec["vpd"]].mean(),
                    "VPD_factor_min": active_group[spec["vpd"]].min(),
                    "Water_stress_factor_mean": active_group[
                        spec["water"]
                    ].mean(),
                    "Water_stress_factor_min": active_group[
                        spec["water"]
                    ].min(),
                    "Soil_water_mean": active_group[
                        "TotalSoilWaterasTotalSoilWater"
                    ].mean(),
                    "Soil_water_min": active_group[
                        "TotalSoilWaterasTotalSoilWater"
                    ].min(),
                    "Soil_N_mean": active_group["TotalSoilN"].mean(),
                    "Soil_N_min": active_group["TotalSoilN"].min(),
                    "Nitrogen_applied_kg_ha": group[
                        "NitrogenApplied"
                    ].sum(min_count=1),
                    "Soil_evaporation_active_sum": active_group[
                        "SoilEvaporation"
                    ].sum(min_count=1),
                    "PAWC_mean": active_group["PAWC"].mean(),
                    "OM_residue_mean": active_group["OM_residue"].mean(),
                    "SOC_change_t_C_ha_yr": soc_end - soc_start,
                    "Active_days": int(active.sum()),
                    "Harvest_DOY": int(harvest_date.dayofyear),
                }
            )

    return pd.DataFrame(annual_soc_rows), pd.DataFrame(crop_rows)


def read_individual_csv(path: Path) -> pd.DataFrame:
    header = pd.read_csv(path, nrows=0).columns.tolist()
    usecols = [col for col in SELECTED_COLUMNS if col in header]
    return pd.read_csv(
        path,
        usecols=usecols,
        low_memory=False,
        encoding="utf-8-sig",
    )


def read_combined_xlsx(path: Path) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["DailyReport"]
    iterator = worksheet.iter_rows(values_only=True)
    header = list(next(iterator))
    index = {column: i for i, column in enumerate(header)}
    available = [column for column in SELECTED_COLUMNS if column in index]
    rows = [
        {column: row[index[column]] for column in available}
        for row in iterator
    ]
    workbook.close()
    return pd.DataFrame.from_records(rows)


def period_from_filename(path: Path) -> str:
    return "Future" if "Future" in path.name else "Historical"


def combined_system_from_filename(path: Path) -> str:
    text = path.name.lower()
    if "ideal" in text:
        return "Ideal farming system"
    if "most_prod" in text:
        return "Most productive"
    if "least_prod" in text:
        return "Least productive"
    raise ValueError(f"Cannot classify combined-system file: {path.name}")


def safe_stem(path: Path) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", path.stem).strip("_")


def process_source(
    path: Path,
    *,
    scope: str,
    period: str | None = None,
    system: str | None = None,
) -> None:
    resolved_period = period or period_from_filename(path)
    resolved_system = system
    if scope == "Combined" and resolved_system is None:
        resolved_system = combined_system_from_filename(path)

    print(f"Reading {scope.lower()} file: {path.name}", flush=True)
    data = (
        read_individual_csv(path)
        if scope == "Individual"
        else read_combined_xlsx(path)
    )
    soc, crop_yield = summarize_one_source(
        data,
        period=resolved_period,
        scope=scope,
        system_override=resolved_system,
    )

    stem = safe_stem(path)
    soc_path = INTERMEDIATE_DIR / f"{stem}_annual_soc.csv"
    yield_path = INTERMEDIATE_DIR / f"{stem}_annual_yield.csv"
    soc.to_csv(soc_path, index=False)
    crop_yield.to_csv(yield_path, index=False)
    print(
        f"Wrote {len(soc):,} SOC and {len(crop_yield):,} crop-year "
        f"records for {path.name}",
        flush=True,
    )

    del data, soc, crop_yield
    gc.collect()


def finalize_outputs() -> None:
    soc_paths = sorted(INTERMEDIATE_DIR.glob("*_annual_soc.csv"))
    yield_paths = sorted(INTERMEDIATE_DIR.glob("*_annual_yield.csv"))
    if not soc_paths or not yield_paths:
        raise FileNotFoundError(
            "No intermediate annual files were found. Run --source first."
        )

    annual_soc = pd.concat(
        [pd.read_csv(path) for path in soc_paths],
        ignore_index=True,
        sort=False,
    )
    annual_yield = pd.concat(
        [pd.read_csv(path) for path in yield_paths],
        ignore_index=True,
        sort=False,
    )

    period_start = {"Historical": 1993, "Future": 2030}
    annual_soc["Year_index"] = annual_soc.apply(
        lambda row: row["Year"] - period_start[row["Period"]] + 1,
        axis=1,
    )
    annual_yield["Year_index"] = annual_yield.apply(
        lambda row: row["Year"] - period_start[row["Period"]] + 1,
        axis=1,
    )

    annual_soc["OM_residue_change"] = (
        annual_soc["OM_residue_end"] - annual_soc["OM_residue_start"]
    )

    sort_columns = [
        "Scope",
        "System",
        "Period",
        "CO2_scenario",
        "Rotation",
        "Year",
    ]
    annual_soc = annual_soc.sort_values(sort_columns).reset_index(drop=True)
    annual_yield = annual_yield.sort_values(
        sort_columns + ["Crop"]
    ).reset_index(drop=True)

    annual_soc.to_csv(
        OUTPUT_DIR / "RA_annual_SOC_process_data.csv",
        index=False,
    )
    annual_yield.to_csv(
        OUTPUT_DIR / "RA_annual_crop_yield_process_data.csv",
        index=False,
    )

    coverage_soc = (
        annual_soc.groupby(
            [
                "Scope",
                "System",
                "Period",
                "CO2_scenario",
                "Rotation",
            ],
            dropna=False,
            observed=True,
        )
        .agg(
            n_years=("Year", "nunique"),
            first_year=("Year", "min"),
            last_year=("Year", "max"),
            n_records=("Year", "size"),
            mean_daily_rows=("n_daily_rows", "mean"),
            missing_SOC_change=("SOC_change_t_C_ha_yr", lambda x: x.isna().sum()),
        )
        .reset_index()
    )

    coverage_yield = (
        annual_yield.groupby(
            [
                "Scope",
                "System",
                "Period",
                "CO2_scenario",
                "Rotation",
                "Crop",
            ],
            dropna=False,
            observed=True,
        )
        .agg(
            n_crop_years=("Year", "nunique"),
            first_year=("Year", "min"),
            last_year=("Year", "max"),
            mean_yield_t_ha=("Yield_t_ha", "mean"),
        )
        .reset_index()
    )

    coverage_soc.to_csv(
        OUTPUT_DIR / "RA_SOC_coverage_audit.csv",
        index=False,
    )
    coverage_yield.to_csv(
        OUTPUT_DIR / "RA_yield_coverage_audit.csv",
        index=False,
    )

    print(
        f"Wrote {len(annual_soc):,} annual SOC records and "
        f"{len(annual_yield):,} crop-year records to {OUTPUT_DIR}",
        flush=True,
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path)
    parser.add_argument(
        "--scope",
        choices=["Individual", "Combined"],
    )
    parser.add_argument(
        "--period",
        choices=["Historical", "Future"],
    )
    parser.add_argument("--system")
    parser.add_argument("--finalize", action="store_true")
    args = parser.parse_args()

    if args.finalize:
        finalize_outputs()
        return

    if args.source is None or args.scope is None:
        parser.error("--source and --scope are required unless --finalize is used")

    process_source(
        args.source,
        scope=args.scope,
        period=args.period,
        system=args.system,
    )


if __name__ == "__main__":
    main()
